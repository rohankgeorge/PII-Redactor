"""Tests for split analyze/apply redaction API endpoints."""

from io import BytesIO

from docx import Document as DocxDocument
from fastapi.testclient import TestClient

import server


def _fake_process_document(_content: bytes):
    doc = DocxDocument()
    doc.add_paragraph("Person 1")
    out = BytesIO()
    doc.save(out)
    return (
        {"INDIVIDUAL": 1},
        1,
        out.getvalue(),
        [
            {
                "category": "INDIVIDUAL",
                "placeholder": "Person 1",
                "location": "Paragraph 1",
                "original_text": "Alice",
            }
        ],
        [],
    )


def test_analyze_stages_document_and_returns_analysis_id(monkeypatch):
    monkeypatch.setattr(server.nlp_engine, "initialize_models", lambda: None)
    monkeypatch.setattr(server, "process_document", _fake_process_document)

    client = TestClient(server.app)
    response = client.post(
        "/api/analyze",
        files={
            "file": (
                "sample.docx",
                b"fake-docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["analysis_id"]
    assert payload["total"] == 1
    assert payload["review"]["qa_signal_count"] == 0


def test_apply_redaction_rejects_unknown_analysis_id(monkeypatch):
    monkeypatch.setattr(server.nlp_engine, "initialize_models", lambda: None)

    client = TestClient(server.app)
    response = client.post("/api/apply-redaction", json={"analysis_id": "missing-id"})

    assert response.status_code == 404
    assert "re-analyze" in response.json()["detail"].lower()


def test_apply_redaction_from_analysis_returns_downloadable_result(monkeypatch):
    monkeypatch.setattr(server.nlp_engine, "initialize_models", lambda: None)
    monkeypatch.setattr(server, "process_document", _fake_process_document)

    client = TestClient(server.app)
    analyzed = client.post(
        "/api/analyze",
        files={
            "file": (
                "contract.docx",
                b"fake-docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert analyzed.status_code == 200
    analysis_id = analyzed.json()["analysis_id"]

    applied = client.post("/api/apply-redaction", json={"analysis_id": analysis_id})
    assert applied.status_code == 200

    payload = applied.json()
    assert payload["analysis_id"] == analysis_id
    assert payload["file_id"]
    assert payload["filename"].startswith("redacted_contract")
    assert payload["total"] == 1


def test_analyze_returns_candidates(monkeypatch):
    monkeypatch.setattr(server.nlp_engine, "initialize_models", lambda: None)
    monkeypatch.setattr(server, "process_document", _fake_process_document)

    client = TestClient(server.app)
    response = client.post(
        "/api/analyze",
        files={
            "file": (
                "candidate.docx",
                b"fake-docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert len(payload["candidates"]) == 1
    assert payload["candidates"][0]["category"] == "INDIVIDUAL"
    assert payload["candidates"][0]["candidate_id"].startswith("cand_")


def test_apply_redaction_exclude_candidate_restores_original(monkeypatch):
    monkeypatch.setattr(server.nlp_engine, "initialize_models", lambda: None)
    monkeypatch.setattr(server, "process_document", _fake_process_document)

    client = TestClient(server.app)
    analyzed = client.post(
        "/api/analyze",
        files={
            "file": (
                "filter.docx",
                b"fake-docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert analyzed.status_code == 200
    analysis_payload = analyzed.json()
    candidate_id = analysis_payload["candidates"][0]["candidate_id"]

    applied = client.post(
        "/api/apply-redaction",
        json={
            "analysis_id": analysis_payload["analysis_id"],
            "exclude_candidate_ids": [candidate_id],
        },
    )
    assert applied.status_code == 200

    payload = applied.json()
    assert payload["total"] == 0
    assert payload["stats"] == {}
    assert payload["audit_log"] == []


def test_apply_redaction_rejects_unknown_candidate_id(monkeypatch):
    monkeypatch.setattr(server.nlp_engine, "initialize_models", lambda: None)
    monkeypatch.setattr(server, "process_document", _fake_process_document)

    client = TestClient(server.app)
    analyzed = client.post(
        "/api/analyze",
        files={
            "file": (
                "unknown-id.docx",
                b"fake-docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert analyzed.status_code == 200
    analysis_id = analyzed.json()["analysis_id"]

    applied = client.post(
        "/api/apply-redaction",
        json={"analysis_id": analysis_id, "exclude_candidate_ids": ["cand_99999"]},
    )
    assert applied.status_code == 400
    assert "unknown candidate id" in applied.json()["detail"].lower()


def test_apply_redaction_rejects_mixed_include_and_exclude_filters(monkeypatch):
    monkeypatch.setattr(server.nlp_engine, "initialize_models", lambda: None)
    monkeypatch.setattr(server, "process_document", _fake_process_document)

    client = TestClient(server.app)
    analyzed = client.post(
        "/api/analyze",
        files={
            "file": (
                "mixed-filters.docx",
                b"fake-docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert analyzed.status_code == 200
    analysis_payload = analyzed.json()
    candidate_id = analysis_payload["candidates"][0]["candidate_id"]

    applied = client.post(
        "/api/apply-redaction",
        json={
            "analysis_id": analysis_payload["analysis_id"],
            "include_candidate_ids": [candidate_id],
            "exclude_candidate_ids": [candidate_id],
        },
    )
    assert applied.status_code == 422


def test_redact_endpoint_remains_compatible(monkeypatch):
    monkeypatch.setattr(server.nlp_engine, "initialize_models", lambda: None)
    monkeypatch.setattr(server, "process_document", _fake_process_document)

    client = TestClient(server.app)
    response = client.post(
        "/api/redact",
        files={
            "file": (
                "compat.docx",
                b"fake-docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["file_id"]
    assert payload["filename"].startswith("redacted_compat")
    assert payload["stats"] == {"INDIVIDUAL": 1}
    assert payload["total"] == 1


def _fake_process_document_with_location(_content: bytes):
    doc = DocxDocument()
    doc.add_paragraph("[Location 1] Person 1")
    out = BytesIO()
    doc.save(out)
    return (
        {"LOCATION": 1, "INDIVIDUAL": 1},
        2,
        out.getvalue(),
        [
            {
                "category": "LOCATION",
                "placeholder": "[Location 1]",
                "location": "Paragraph 1",
                "original_text": "India",
            },
            {
                "category": "INDIVIDUAL",
                "placeholder": "Person 1",
                "location": "Paragraph 1",
                "original_text": "Alice",
            },
        ],
        [],
    )


def _fake_process_document_with_city_and_country(_content: bytes):
    doc = DocxDocument()
    doc.add_paragraph("[Location 1] [Location 2] Person 1")
    out = BytesIO()
    doc.save(out)
    return (
        {"LOCATION": 2, "INDIVIDUAL": 1},
        3,
        out.getvalue(),
        [
            {
                "category": "LOCATION",
                "placeholder": "[Location 1]",
                "location": "Paragraph 1",
                "original_text": "India",
            },
            {
                "category": "LOCATION",
                "placeholder": "[Location 2]",
                "location": "Paragraph 1",
                "original_text": "Mumbai",
            },
            {
                "category": "INDIVIDUAL",
                "placeholder": "Person 1",
                "location": "Paragraph 1",
                "original_text": "Alice",
            },
        ],
        [],
    )


def test_analyze_candidates_include_policy_tags(monkeypatch):
    monkeypatch.setattr(server.nlp_engine, "initialize_models", lambda: None)
    monkeypatch.setattr(server, "process_document", _fake_process_document_with_location)

    client = TestClient(server.app)
    analyzed = client.post(
        "/api/analyze",
        files={
            "file": (
                "policy.docx",
                b"fake-docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert analyzed.status_code == 200

    candidates = analyzed.json()["candidates"]
    location_candidate = next(candidate for candidate in candidates if candidate["category"] == "LOCATION")
    assert set(location_candidate["policy_tags"]) == {"country", "location"}


def test_apply_redaction_policy_toggle_excludes_locations(monkeypatch):
    monkeypatch.setattr(server.nlp_engine, "initialize_models", lambda: None)
    monkeypatch.setattr(server, "process_document", _fake_process_document_with_location)

    client = TestClient(server.app)
    analyzed = client.post(
        "/api/analyze",
        files={
            "file": (
                "policy.docx",
                b"fake-docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert analyzed.status_code == 200
    analysis_id = analyzed.json()["analysis_id"]

    applied = client.post(
        "/api/apply-redaction",
        json={
            "analysis_id": analysis_id,
            "redact_locations": False,
            "redact_countries": False,
        },
    )
    assert applied.status_code == 200

    payload = applied.json()
    assert payload["total"] == 1
    assert payload["stats"] == {"INDIVIDUAL": 1}
    assert len(payload["audit_log"]) == 1
    assert payload["audit_log"][0]["category"] == "INDIVIDUAL"


def test_apply_redaction_country_toggle_only_keeps_city_redacted(monkeypatch):
    monkeypatch.setattr(server.nlp_engine, "initialize_models", lambda: None)
    monkeypatch.setattr(server, "process_document", _fake_process_document_with_city_and_country)

    client = TestClient(server.app)
    analyzed = client.post(
        "/api/analyze",
        files={
            "file": (
                "country-policy.docx",
                b"fake-docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert analyzed.status_code == 200
    analysis_id = analyzed.json()["analysis_id"]

    applied = client.post(
        "/api/apply-redaction",
        json={
            "analysis_id": analysis_id,
            "redact_countries": False,
        },
    )
    assert applied.status_code == 200

    payload = applied.json()
    assert payload["total"] == 2
    assert payload["stats"] == {"INDIVIDUAL": 1, "LOCATION": 1}
    location_rows = [row for row in payload["audit_log"] if row["category"] == "LOCATION"]
    assert len(location_rows) == 1
    assert location_rows[0]["original_text"] == "Mumbai"
