from fastapi import FastAPI, APIRouter, UploadFile, File, HTTPException
from fastapi.responses import Response
from fastapi.staticfiles import StaticFiles
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
import re
import io
import csv
import uuid
import time
import base64
import subprocess
import tempfile
from pathlib import Path
from typing import Dict, List, Optional
from pydantic import BaseModel, Field, model_validator
from docx import Document as DocxDocument
from pypdf import PdfReader

from indian_pii_data import (
    PII_REGEX_PATTERNS,
    INDIAN_FIRST_NAMES,
    INDIAN_SURNAMES,
    INDIAN_CITIES,
    INDIAN_STATES,
)
from pii_engine import PIITracker, redact_text
import nlp_engine
import format_converter
import rule_library

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")

# ── Writable data directory for mutable rule data ──────────
import sys as _sys
if getattr(_sys, "_MEIPASS", None):
    DATA_DIR = Path(os.environ.get("APPDATA", str(ROOT_DIR))) / "PII Redactor" / "data"
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    # Bootstrap: copy bundled seed files to writable location on first run
    for _fname in ("user_allow_list.txt", "user_redact_list.txt"):
        _dest = DATA_DIR / _fname
        _src = ROOT_DIR / _fname
        if not _dest.exists() and _src.exists():
            import shutil
            shutil.copy2(str(_src), str(_dest))
else:
    DATA_DIR = ROOT_DIR

# MongoDB connection (optional for local runs)
mongo_url = os.environ.get("MONGO_URL", "mongodb://localhost:27017")
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ.get("DB_NAME", "pii_redactor")]

app = FastAPI()
api_router = APIRouter(prefix="/api")

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
)
if getattr(_sys, "_MEIPASS", None):
    _log_dir = Path(os.environ.get("APPDATA", str(ROOT_DIR))) / "PII Redactor" / "logs"
    _log_dir.mkdir(parents=True, exist_ok=True)
    _fh = logging.FileHandler(str(_log_dir / "backend.log"), encoding="utf-8")
    _fh.setFormatter(logging.Formatter("%(asctime)s - %(name)s - %(levelname)s - %(message)s"))
    logging.getLogger().addHandler(_fh)
logger = logging.getLogger(__name__)
if getattr(_sys, "_MEIPASS", None):
    logger.info("PyInstaller bundle detected — logging to %s", _log_dir / "backend.log")


class RuleCreateRequest(BaseModel):
    term: str
    mode: str
    enabled: bool = True


class RuleUpdateRequest(BaseModel):
    term: Optional[str] = None
    mode: Optional[str] = None
    enabled: Optional[bool] = None


class ApplyRedactionRequest(BaseModel):
    analysis_id: str
    include_candidate_ids: Optional[List[str]] = Field(default=None)
    exclude_candidate_ids: Optional[List[str]] = Field(default=None)
    redact_locations: bool = True
    redact_countries: bool = True
    force_redact_terms: Optional[List[str]] = Field(default=None)
    placeholder_overrides: Optional[Dict[str, str]] = Field(default=None)

    @model_validator(mode="after")
    def validate_filters(self):
        if self.include_candidate_ids and self.exclude_candidate_ids:
            raise ValueError("Provide either include_candidate_ids or exclude_candidate_ids, not both")
        return self


MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB

# ── Temporary in-memory file store for reliable downloads ────
_file_store: Dict[str, dict] = {}
_analysis_store: Dict[str, dict] = {}
_FILE_TTL = 600  # 10 minutes


def _store_file(
    data: bytes,
    filename: str,
    audit_log: list,
    stats: dict,
    total: int,
    qa_signals: list,
) -> str:
    """Store file data and return a unique download ID."""
    _cleanup_expired()
    file_id = uuid.uuid4().hex
    _file_store[file_id] = {
        "data": data,
        "filename": filename,
        "audit_log": audit_log,
        "qa_signals": qa_signals,
        "stats": stats,
        "total": total,
        "created": time.time(),
    }
    return file_id


def _portable_review_payload(audit_log: list, qa_signals: list) -> dict:
    """Return UI-agnostic review hints for any client (web/desktop/offline)."""
    potential_leaks = [
        row for row in audit_log
        if row.get("category") == "POTENTIAL_LEAK"
    ]
    return {
        "review_required": bool(potential_leaks),
        "highlight_color": "RED",
        "potential_leak_count": len(potential_leaks),
        "manual_review_items": potential_leaks,
        "qa_signal_count": len(qa_signals),
    }


def _build_candidates(audit_log: list) -> list:
    """Build redaction candidates from detected audit rows for explicit apply-time filtering."""
    candidates = []
    for index, row in enumerate(audit_log, start=1):
        if row.get("category") == "POTENTIAL_LEAK":
            continue
        placeholder = row.get("placeholder", "")
        original_text = row.get("original_text", "")
        if not placeholder or not original_text:
            continue
        candidate_id = f"cand_{index:05d}"
        policy_tags = _policy_tags_for_row(row)
        candidates.append(
            {
                "candidate_id": candidate_id,
                "category": row.get("category", "UNKNOWN"),
                "placeholder": placeholder,
                "original_text": original_text,
                "location": row.get("location", ""),
                "policy_tags": policy_tags,
            }
        )
    return candidates


_COUNTRY_TERMS = {
    "india",
    "united states",
    "united states of america",
    "usa",
    "uk",
    "united kingdom",
    "england",
    "scotland",
    "wales",
    "ireland",
    "australia",
    "canada",
    "singapore",
    "uae",
    "united arab emirates",
    "germany",
    "france",
    "spain",
    "italy",
    "china",
    "japan",
    "south korea",
    "russia",
    "brazil",
    "south africa",
}


def _policy_tags_for_row(row: dict) -> list[str]:
    """Assign policy tags so apply-redaction can toggle location/country categories."""
    category = str(row.get("category", "")).upper().strip()
    original_text = str(row.get("original_text", "")).strip().casefold()
    tags: set[str] = set()

    if category in {"LOCATION", "ADDRESS"}:
        tags.add("location")

    if category == "LOCATION" and original_text in _COUNTRY_TERMS:
        tags.add("country")

    return sorted(tags)


def _store_analysis(
    source_bytes: bytes,
    redacted_bytes: bytes,
    source_name: str,
    source_type: str,
    stats: dict,
    total: int,
    audit_log: list,
    qa_signals: list,
    candidates: list,
) -> str:
    """Store analyzed source material for a later apply-redaction step."""
    _cleanup_expired()
    analysis_id = uuid.uuid4().hex
    _analysis_store[analysis_id] = {
        "source_bytes": source_bytes,
        "source_name": source_name,
        "source_type": source_type,
        "redacted_bytes": redacted_bytes,
        "stats": stats,
        "total": total,
        "audit_log": audit_log,
        "qa_signals": qa_signals,
        "candidates": candidates,
        "review": _portable_review_payload(audit_log, qa_signals),
        "created": time.time(),
    }
    return analysis_id


def _apply_placeholder_reverts(redacted_docx_bytes: bytes, rows_to_revert: list[dict]) -> bytes:
    """Restore selected placeholders back to original text inside a DOCX artifact."""
    if not rows_to_revert:
        return redacted_docx_bytes

    replacement_map = {
        row["placeholder"]: row.get("original_text", "")
        for row in rows_to_revert
        if row.get("placeholder") and row.get("original_text")
    }
    if not replacement_map:
        return redacted_docx_bytes

    doc = DocxDocument(io.BytesIO(redacted_docx_bytes))

    def replace_text(value: str) -> str:
        updated = value
        for placeholder, original_text in replacement_map.items():
            updated = updated.replace(placeholder, original_text)
        return updated

    def rewrite_paragraph(para):
        if not para.text:
            return
        new_text = replace_text(para.text)
        if new_text == para.text:
            return
        if para.runs:
            for index, run in enumerate(para.runs):
                run.text = new_text if index == 0 else ""
        else:
            para.add_run(new_text)

    for para in doc.paragraphs:
        rewrite_paragraph(para)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    rewrite_paragraph(para)

    for section in doc.sections:
        for para in section.header.paragraphs:
            rewrite_paragraph(para)
        for para in section.footer.paragraphs:
            rewrite_paragraph(para)

    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out.read()


def _apply_force_redact_terms(docx_bytes: bytes, terms: list[str], audit_log: list, stats: dict, total: int) -> tuple[bytes, dict, int, list]:
    """Find and redact additional user-specified terms inside a DOCX artifact."""
    if not terms:
        return docx_bytes, stats, total, audit_log

    tracker = PIITracker()
    doc = DocxDocument(io.BytesIO(docx_bytes))
    new_audit_rows: list[dict] = []

    def _replace_in_text(text: str, context: str) -> str:
        for term in sorted(terms, key=len, reverse=True):
            if not term.strip():
                continue
            pattern = re.compile(re.escape(term), re.IGNORECASE)
            matches = list(pattern.finditer(text))
            for m in reversed(matches):
                placeholder = tracker.placeholder("INDIVIDUAL", m.group(), context)
                new_audit_rows.append({
                    "category": "INDIVIDUAL",
                    "placeholder": placeholder,
                    "original_text": m.group(),
                    "location": context,
                })
                text = text[:m.start()] + placeholder + text[m.end():]
        return text

    def _rewrite_paragraph(para, context: str):
        if not para.text:
            return
        new_text = _replace_in_text(para.text, context)
        if new_text == para.text:
            return
        if para.runs:
            for idx, run in enumerate(para.runs):
                run.text = new_text if idx == 0 else ""
        else:
            para.add_run(new_text)

    for i, para in enumerate(doc.paragraphs):
        _rewrite_paragraph(para, f"Paragraph {i + 1}")
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _rewrite_paragraph(para, "Table cell")

    if not new_audit_rows:
        return docx_bytes, stats, total, audit_log

    updated_stats = dict(stats)
    for row in new_audit_rows:
        cat = row["category"]
        updated_stats[cat] = updated_stats.get(cat, 0) + 1
    updated_total = total + len(new_audit_rows)
    updated_audit = audit_log + new_audit_rows

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read(), updated_stats, updated_total, updated_audit


def _apply_placeholder_overrides(docx_bytes: bytes, overrides: dict[str, str], candidates: list, audit_log: list) -> tuple[bytes, list]:
    """Replace auto-generated placeholders with user-specified custom text in a DOCX."""
    if not overrides:
        return docx_bytes, audit_log

    candidate_lookup = {c["candidate_id"]: c for c in candidates}
    replacement_map: dict[str, str] = {}
    for cand_id, custom_text in overrides.items():
        cand = candidate_lookup.get(cand_id)
        if cand and cand.get("placeholder"):
            replacement_map[cand["placeholder"]] = custom_text

    if not replacement_map:
        return docx_bytes, audit_log

    doc = DocxDocument(io.BytesIO(docx_bytes))

    def _replace_text(value: str) -> str:
        updated = value
        for old_ph, new_ph in replacement_map.items():
            updated = updated.replace(old_ph, new_ph)
        return updated

    def _rewrite_paragraph(para):
        if not para.text:
            return
        new_text = _replace_text(para.text)
        if new_text == para.text:
            return
        if para.runs:
            for idx, run in enumerate(para.runs):
                run.text = new_text if idx == 0 else ""
        else:
            para.add_run(new_text)

    for para in doc.paragraphs:
        _rewrite_paragraph(para)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _rewrite_paragraph(para)
    for section in doc.sections:
        for para in section.header.paragraphs:
            _rewrite_paragraph(para)
        for para in section.footer.paragraphs:
            _rewrite_paragraph(para)

    updated_audit = []
    for row in audit_log:
        ph = row.get("placeholder", "")
        if ph in replacement_map:
            updated_audit.append({**row, "placeholder": replacement_map[ph]})
        else:
            updated_audit.append(row)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read(), updated_audit


def _filter_analysis_for_apply(analysis_entry: dict, payload: ApplyRedactionRequest) -> tuple[bytes, dict, int, list]:
    """Apply candidate include/exclude filters and return updated output material."""
    candidates = analysis_entry.get("candidates", [])
    candidate_lookup = {candidate["candidate_id"]: candidate for candidate in candidates}
    include_ids = set(payload.include_candidate_ids or [])
    exclude_ids = set(payload.exclude_candidate_ids or [])

    policy_excluded_ids: set[str] = set()
    for candidate in candidates:
        tags = set(candidate.get("policy_tags") or [])
        if not payload.redact_locations and "location" in tags:
            policy_excluded_ids.add(candidate["candidate_id"])
        if not payload.redact_countries and "country" in tags:
            policy_excluded_ids.add(candidate["candidate_id"])

    if include_ids and not include_ids.issubset(candidate_lookup.keys()):
        raise HTTPException(status_code=400, detail="Unknown candidate id in include_candidate_ids")
    if exclude_ids and not exclude_ids.issubset(candidate_lookup.keys()):
        raise HTTPException(status_code=400, detail="Unknown candidate id in exclude_candidate_ids")

    if include_ids:
        deselected_ids = (set(candidate_lookup.keys()) - include_ids) | (policy_excluded_ids - include_ids)
    else:
        deselected_ids = exclude_ids | policy_excluded_ids

    if not deselected_ids:
        return (
            analysis_entry["redacted_bytes"],
            analysis_entry["stats"],
            analysis_entry["total"],
            analysis_entry["audit_log"],
        )

    deselected_placeholders = {
        candidate_lookup[candidate_id]["placeholder"]
        for candidate_id in deselected_ids
    }
    rows_to_revert = [
        row
        for row in analysis_entry["audit_log"]
        if row.get("placeholder") in deselected_placeholders
    ]

    filtered_audit_log = [
        row
        for row in analysis_entry["audit_log"]
        if row.get("placeholder") not in deselected_placeholders
    ]
    filtered_stats = {}
    filtered_total = 0
    for row in filtered_audit_log:
        category = row.get("category")
        if not category or category == "POTENTIAL_LEAK":
            continue
        filtered_stats[category] = filtered_stats.get(category, 0) + 1
        filtered_total += 1

    filtered_bytes = _apply_placeholder_reverts(analysis_entry["redacted_bytes"], rows_to_revert)
    return filtered_bytes, filtered_stats, filtered_total, filtered_audit_log


def _cleanup_expired():
    now = time.time()
    expired = [k for k, v in _file_store.items() if now - v["created"] > _FILE_TTL]
    for k in expired:
        del _file_store[k]
    expired_analysis = [k for k, v in _analysis_store.items() if now - v["created"] > _FILE_TTL]
    for k in expired_analysis:
        del _analysis_store[k]


async def _validate_upload_and_read(file: UploadFile) -> tuple[str, bytes]:
    """Validate supported file extension and upload size, then read bytes."""
    fname = (file.filename or "").lower()
    if not (fname.endswith(".docx") or fname.endswith(".doc") or fname.endswith(".pdf")):
        raise HTTPException(status_code=400, detail="Only .doc, .docx, and text-based .pdf files are supported")

    content = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=400, detail="File too large. Maximum size is 10 MB.")
    return fname, content


def _run_redaction(content: bytes, fname: str):
    """Run the redaction engine for DOC/DOCX/PDF bytes and return pipeline outputs."""
    if fname.endswith(".pdf"):
        return process_pdf_document(content)

    # Convert legacy .doc → .docx when needed
    if fname.endswith(".doc") and not fname.endswith(".docx"):
        content = convert_doc_to_docx(content)

    return process_document(content)


def _build_redact_response(source_filename: str, redacted_bytes: bytes, stats: dict, total: int, audit_log: list, qa_signals: list):
    """Store redacted bytes and return standard API response payload."""
    original_stem = (source_filename or "document").rsplit(".", 1)[0]
    filename = f"redacted_{original_stem}.docx"

    file_id = _store_file(redacted_bytes, filename, audit_log, stats, total, qa_signals)

    return {
        "stats": stats,
        "total": total,
        "qa_signals": qa_signals,
        "file_id": file_id,
        "filename": filename,
        "audit_log": audit_log,
        "review": _portable_review_payload(audit_log, qa_signals),
    }


# ── .doc → .docx conversion ─────────────────────────────────
def convert_doc_to_docx(doc_bytes: bytes) -> bytes:
    """Convert legacy .doc to .docx via antiword text extraction."""
    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as tmp:
            tmp.write(doc_bytes)
            tmp_path = tmp.name

        env = {**os.environ, "HOME": "/tmp"}
        try:
            result = subprocess.run(
                ["antiword", tmp_path],
                capture_output=True,
                text=True,
                timeout=30,
                env=env,
            )
            if result.returncode != 0:
                logger.warning("antiword failed, attempting plain-text fallback: %s", result.stderr.strip())
                text = doc_bytes.decode("utf-8", errors="ignore")
            else:
                text = result.stdout
        except FileNotFoundError:
            logger.warning("antiword executable not found; using plain-text fallback for .doc conversion")
            text = doc_bytes.decode("utf-8", errors="ignore")

        doc = DocxDocument()
        for line in text.split("\n"):
            doc.add_paragraph(line)

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf.read()
    finally:
        if tmp_path and os.path.exists(tmp_path):
            os.unlink(tmp_path)


# ── Core redaction logic ────────────────────────────────────
def _process_paragraph(para, tracker: PIITracker, ctx: str):
    full_text = para.text
    if not full_text.strip():
        return
    new_text = redact_text(full_text, tracker, ctx)
    if new_text != full_text and para.runs:
        for i, run in enumerate(para.runs):
            run.text = new_text if i == 0 else ""


def process_document(doc_bytes: bytes):
    doc = DocxDocument(io.BytesIO(doc_bytes))
    tracker = PIITracker()
    para_num = 0

    # Paragraphs
    for para in doc.paragraphs:
        para_num += 1
        _process_paragraph(para, tracker, f"Paragraph {para_num}")

    # Tables (including nested)
    table_num = 0

    def _process_table(table, prefix=""):
        nonlocal table_num
        table_num += 1
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                ctx = f"{prefix}Table {table_num}, Row {r_idx+1}, Col {c_idx+1}"
                for para in cell.paragraphs:
                    _process_paragraph(para, tracker, ctx)
                for nested in cell.tables:
                    _process_table(nested, prefix=f"{ctx} > ")

    for table in doc.tables:
        _process_table(table)

    # Headers & footers
    for s_idx, section in enumerate(doc.sections):
        if section.header:
            for para in section.header.paragraphs:
                _process_paragraph(para, tracker, f"Header (section {s_idx+1})")
        if section.footer:
            for para in section.footer.paragraphs:
                _process_paragraph(para, tracker, f"Footer (section {s_idx+1})")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return tracker.stats, tracker.total, buf.read(), tracker.audit_log, tracker.qa_signals


def process_pdf_document(pdf_bytes: bytes):
    """Extract text from a PDF with text layers, redact it, and return DOCX bytes."""
    reader = PdfReader(io.BytesIO(pdf_bytes))
    tracker = PIITracker()
    output_doc = DocxDocument()
    has_extractable_text = False

    for page_num, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        if not page_text.strip():
            continue

        has_extractable_text = True
        output_doc.add_paragraph(f"--- Page {page_num} ---")

        # Redact page text as a whole to preserve multi-line context.
        redacted_page_text = redact_text(page_text, tracker, f"Page {page_num}")
        for line in redacted_page_text.splitlines():
            output_doc.add_paragraph(line)

    if not has_extractable_text:
        raise ValueError("PDF has no extractable text layer (image-based PDFs are not supported yet).")

    buf = io.BytesIO()
    output_doc.save(buf)
    buf.seek(0)
    return tracker.stats, tracker.total, buf.read(), tracker.audit_log, tracker.qa_signals


# ── Routes ──────────────────────────────────────────────────
@api_router.get("/")
async def root():
    return {"message": "RedactAI API is running"}


@api_router.post("/redact")
async def redact_document(file: UploadFile = File(...)):
    try:
        fname, content = await _validate_upload_and_read(file)
        stats, total, redacted_bytes, audit_log, qa_signals = _run_redaction(content, fname)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    except Exception as exc:
        logger.error("Document processing failed: %s", exc, exc_info=True)
        raise HTTPException(status_code=500, detail=f"Processing failed: {exc}")

    return _build_redact_response(file.filename or "document", redacted_bytes, stats, total, audit_log, qa_signals)


@api_router.post("/analyze")
async def analyze_document(file: UploadFile = File(...)):
    """Analyze an uploaded file and stage it for explicit apply-redaction."""
    try:
        fname, content = await _validate_upload_and_read(file)
        stats, total, redacted_bytes, audit_log, qa_signals = _run_redaction(content, fname)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    except Exception as exc:
        logger.error("Document analysis failed: %s", exc, exc_info=True)
        raise HTTPException(status_code=500, detail=f"Analysis failed: {exc}")

    candidates = _build_candidates(audit_log)
    analysis_id = _store_analysis(
        content,
        redacted_bytes,
        file.filename or "document",
        fname,
        stats,
        total,
        audit_log,
        qa_signals,
        candidates,
    )
    return {
        "analysis_id": analysis_id,
        "source_filename": file.filename or "document",
        "stats": stats,
        "total": total,
        "qa_signals": qa_signals,
        "audit_log": audit_log,
        "candidates": candidates,
        "review": _portable_review_payload(audit_log, qa_signals),
    }


@api_router.post("/apply-redaction")
async def apply_redaction(payload: ApplyRedactionRequest):
    """Apply redaction to a previously analyzed document and return downloadable artifact."""
    _cleanup_expired()
    analysis_entry = _analysis_store.get(payload.analysis_id)
    if not analysis_entry:
        raise HTTPException(status_code=404, detail="Analysis expired or not found. Please re-analyze.")

    try:
        redacted_bytes, stats, total, audit_log = _filter_analysis_for_apply(analysis_entry, payload)
        redacted_bytes, stats, total, audit_log = _apply_force_redact_terms(
            redacted_bytes, payload.force_redact_terms or [], audit_log, stats, total,
        )
        candidates = analysis_entry.get("candidates", [])
        redacted_bytes, audit_log = _apply_placeholder_overrides(
            redacted_bytes, payload.placeholder_overrides or {}, candidates, audit_log,
        )
    except HTTPException:
        raise
    except Exception as exc:
        logger.error("Document apply-redaction failed: %s", exc, exc_info=True)
        raise HTTPException(status_code=500, detail="Apply redaction failed")

    source_name = analysis_entry["source_name"]
    qa_signals = analysis_entry.get("qa_signals", [])
    response_payload = _build_redact_response(source_name, redacted_bytes, stats, total, audit_log, qa_signals)
    response_payload["analysis_id"] = payload.analysis_id
    return response_payload


@api_router.get("/rules")
async def list_rules():
    return {"rules": rule_library.list_rules()}


@api_router.post("/rules")
async def create_rule(payload: RuleCreateRequest):
    try:
        rule = rule_library.create_rule(payload.term, payload.mode, payload.enabled)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return {"rule": rule}


@api_router.put("/rules/{rule_id}")
async def update_rule(rule_id: str, payload: RuleUpdateRequest):
    try:
        rule = rule_library.update_rule(rule_id, payload.term, payload.mode, payload.enabled)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return {"rule": rule}


@api_router.delete("/rules/{rule_id}")
async def delete_rule(rule_id: str):
    deleted = rule_library.delete_rule(rule_id)
    if not deleted:
        raise HTTPException(status_code=404, detail="rule not found")
    return {"deleted": True, "rule_id": rule_id}


@api_router.get("/download/{file_id}")
async def download_redacted(file_id: str, output_format: str = "docx"):
    entry = _file_store.get(file_id)
    if not entry:
        raise HTTPException(status_code=404, detail="File expired or not found. Please re-process.")

    fmt = output_format.lower().strip().lstrip(".")
    base_stem = entry["filename"].rsplit(".", 1)[0]

    if fmt == "pdf":
        content = format_converter.convert_to_pdf(entry["data"])
        media_type = "application/pdf"
        filename = f"{base_stem}.pdf"
    elif fmt in ("md", "markdown"):
        content = format_converter.convert_to_markdown(entry["data"]).encode("utf-8")
        media_type = "text/markdown; charset=utf-8"
        filename = f"{base_stem}.md"
    else:
        content = entry["data"]
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        filename = entry["filename"]

    return Response(
        content=content,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@api_router.get("/audit-csv/{file_id}")
async def download_audit_csv(file_id: str):
    entry = _file_store.get(file_id)
    if not entry:
        raise HTTPException(status_code=404, detail="File expired or not found. Please re-process.")

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Category", "Placeholder", "Location", "Original Text"])
    for row in entry["audit_log"]:
        writer.writerow([row["category"], row["placeholder"], row["location"], row.get("original_text", "")])
    writer.writerow([])
    writer.writerow(["Summary"])
    writer.writerow(["Category", "Count"])
    for cat, count in sorted(entry["stats"].items(), key=lambda x: -x[1]):
        writer.writerow([cat, count])
    writer.writerow(["TOTAL", entry["total"]])

    csv_bytes = buf.getvalue().encode("utf-8")
    return Response(
        content=csv_bytes,
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="audit_{entry["filename"].replace(".docx", ".csv")}"'},
    )


@api_router.post("/audit-csv-batch")
async def download_batch_audit_csv(file_ids: List[str] = File(default=[])):
    """Generate a single audit CSV combining multiple processed files."""
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Document", "Category", "Placeholder", "Location", "Original Text"])

    for fid in file_ids:
        entry = _file_store.get(fid)
        if not entry:
            continue
        for row in entry["audit_log"]:
            writer.writerow([entry["filename"], row["category"], row["placeholder"], row["location"], row.get("original_text", "")])

    writer.writerow([])
    writer.writerow(["Summary"])
    writer.writerow(["Document", "Total PII", "Status"])
    for fid in file_ids:
        entry = _file_store.get(fid)
        if entry:
            writer.writerow([entry["filename"], entry["total"], "Success"])
        else:
            writer.writerow([fid, 0, "Expired/Not Found"])

    writer.writerow([])
    writer.writerow(["Category Breakdown"])
    writer.writerow(["Document", "Category", "Count"])
    for fid in file_ids:
        entry = _file_store.get(fid)
        if entry:
            for cat, count in sorted(entry["stats"].items(), key=lambda x: -x[1]):
                writer.writerow([entry["filename"], cat, count])

    csv_bytes = buf.getvalue().encode("utf-8")
    return Response(
        content=csv_bytes,
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="pii_audit_report.csv"'},
    )


# ── App wiring ──────────────────────────────────────────────
app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get("CORS_ORIGINS", "*").split(","),
    allow_methods=["*"],
    allow_headers=["*"],
)




# ── Static frontend serving (for standalone/installer builds) ──
_FRONTEND_BUILD_DIR = ROOT_DIR / "frontend_build"
if _FRONTEND_BUILD_DIR.is_dir():
    app.mount("/", StaticFiles(directory=str(_FRONTEND_BUILD_DIR), html=True), name="frontend")


@app.on_event("startup")
async def startup_nlp_models():
    rule_library.initialize(DATA_DIR)
    import asyncio
    loop = asyncio.get_event_loop()
    try:
        await loop.run_in_executor(None, nlp_engine.initialize_models)
        logger.info("NLP models loaded successfully")
    except Exception:
        logger.exception("NLP model loading failed at startup — will retry on first request")


@app.on_event("shutdown")
async def shutdown_db_client():
    if client:
        client.close()


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=8000, log_level="info")
