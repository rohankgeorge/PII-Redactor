"""Tests for audit log original-text tracking and CSV export columns."""

from io import BytesIO

from fastapi.testclient import TestClient
from docx import Document as DocxDocument

import pii_engine
import server


def test_tracker_placeholder_audit_includes_original_text():
    tracker = pii_engine.PIITracker()
    placeholder = tracker.placeholder("EMAIL", "alice@example.com", "Paragraph 1")

    assert placeholder.startswith("[REDACTED_EMAIL")
    assert tracker.audit_log[0]["original_text"] == "alice@example.com"


def test_potential_leak_audit_includes_original_text(monkeypatch):
    tracker = pii_engine.PIITracker()

    monkeypatch.setattr(pii_engine.nlp_engine, "load_nlp_pipeline", object)
    monkeypatch.setattr(
        pii_engine.nlp_engine,
        "detect_names_and_locations",
        lambda text, _nlp: [{"start": 0, "end": 5, "text": "Alice", "label": "PERSON"}],
    )
    monkeypatch.setattr(pii_engine, "QA_AUTO_REDACT_THRESHOLD", 0.99)

    pii_engine._run_final_qa("Alice", tracker, "Paragraph 1")

    leak_rows = [row for row in tracker.audit_log if row.get("category") == "POTENTIAL_LEAK"]
    assert leak_rows
    assert leak_rows[0]["original_text"] == "Alice"


def test_audit_csv_includes_original_text_column(monkeypatch):
    monkeypatch.setattr(server.nlp_engine, "initialize_models", lambda: None)

    def _fake_process_document(_content: bytes):
        doc = DocxDocument()
        out = BytesIO()
        doc.save(out)
        return (
            {"EMAIL": 1},
            1,
            out.getvalue(),
            [{
                "category": "EMAIL",
                "placeholder": "[REDACTED_EMAIL1]",
                "location": "Paragraph 1",
                "original_text": "alice@example.com",
            }],
            [],
        )

    monkeypatch.setattr(server, "process_document", _fake_process_document)

    client = TestClient(server.app)
    response = client.post(
        "/api/redact",
        files={
            "file": (
                "sample.docx",
                b"fake-docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 200
    file_id = response.json()["file_id"]

    csv_response = client.get(f"/api/audit-csv/{file_id}")
    assert csv_response.status_code == 200
    csv_text = csv_response.text
    assert "Category,Placeholder,Location,Original Text" in csv_text
    assert "alice@example.com" in csv_text


def test_batch_audit_csv_includes_original_text_column(monkeypatch):
    monkeypatch.setattr(server.nlp_engine, "initialize_models", lambda: None)

    def _fake_process_document(_content: bytes):
        doc = DocxDocument()
        out = BytesIO()
        doc.save(out)
        return (
            {"PHONE_NUMBER": 1},
            1,
            out.getvalue(),
            [{
                "category": "PHONE_NUMBER",
                "placeholder": "[REDACTED_PHONE_NUMBER1]",
                "location": "Paragraph 2",
                "original_text": "+91 98765 43210",
            }],
            [],
        )

    monkeypatch.setattr(server, "process_document", _fake_process_document)

    client = TestClient(server.app)
    r1 = client.post("/api/redact", files={"file": ("a.docx", b"x", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    r2 = client.post("/api/redact", files={"file": ("b.docx", b"y", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert r1.status_code == 200 and r2.status_code == 200

    batch = client.post(
        "/api/audit-csv-batch",
        files=[("file_ids", (None, r1.json()["file_id"])), ("file_ids", (None, r2.json()["file_id"]))],
    )
    assert batch.status_code == 200
    text = batch.text
    assert "Document,Category,Placeholder,Location,Original Text" in text
    # Phone number starting with + should be escaped with tab
    assert "\t+91 98765 43210" in text


def test_csv_injection_protection_original_text(monkeypatch):
    """Test that formula-like original_text values are escaped to prevent CSV injection."""
    monkeypatch.setattr(server.nlp_engine, "initialize_models", lambda: None)

    def _fake_process_document(_content: bytes):
        doc = DocxDocument()
        out = BytesIO()
        doc.save(out)
        return (
            {"EMAIL": 1},
            1,
            out.getvalue(),
            [{
                "category": "EMAIL",
                "placeholder": "[REDACTED_EMAIL1]",
                "location": "Paragraph 1",
                "original_text": "=1+1@example.com",  # Formula-like email
            }],
            [],
        )

    monkeypatch.setattr(server, "process_document", _fake_process_document)

    client = TestClient(server.app)
    response = client.post(
        "/api/redact",
        files={
            "file": (
                "sample.docx",
                b"fake-docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 200
    file_id = response.json()["file_id"]

    csv_response = client.get(f"/api/audit-csv/{file_id}")
    assert csv_response.status_code == 200
    csv_text = csv_response.text

    # The formula-like value should be escaped with a tab character
    assert "\t=1+1@example.com" in csv_text
    # Should NOT contain the unescaped version at the start of a cell
    lines = csv_text.split('\n')
    for line in lines:
        # Ensure no line starts with the dangerous formula prefix without escape
        if line.startswith("EMAIL,"):
            assert not line.endswith(",=1+1@example.com")


def test_csv_injection_protection_all_columns(monkeypatch):
    """Test that formula prefixes in any column are escaped."""
    monkeypatch.setattr(server.nlp_engine, "initialize_models", lambda: None)

    def _fake_process_document(_content: bytes):
        doc = DocxDocument()
        out = BytesIO()
        doc.save(out)
        return (
            {"+SUSPICIOUS": 1, "-CALC": 1, "@CMD": 1},
            3,
            out.getvalue(),
            [
                {
                    "category": "+SUSPICIOUS",
                    "placeholder": "=evil()",
                    "location": "-location",
                    "original_text": "@cmd",
                },
            ],
            [],
        )

    monkeypatch.setattr(server, "process_document", _fake_process_document)

    client = TestClient(server.app)
    response = client.post(
        "/api/redact",
        files={
            "file": (
                "sample.docx",
                b"fake-docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 200
    file_id = response.json()["file_id"]

    csv_response = client.get(f"/api/audit-csv/{file_id}")
    assert csv_response.status_code == 200
    csv_text = csv_response.text

    # All formula-like values should be tab-escaped
    assert "\t+SUSPICIOUS" in csv_text
    assert "\t=evil()" in csv_text
    assert "\t-location" in csv_text
    assert "\t@cmd" in csv_text
