"""Tests for Milestone 2: Output Format Selection (.docx, .pdf, .md)."""

from io import BytesIO

from docx import Document as DocxDocument
from fastapi.testclient import TestClient

import format_converter
import server


# ── Unit tests for format_converter ──────────────────────────

def _make_docx(*paragraphs, headings=None) -> bytes:
    """Create a minimal DOCX with the given paragraphs and optional headings."""
    doc = DocxDocument()
    if headings:
        for level, text in headings:
            doc.add_heading(text, level=level)
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_convert_to_pdf_returns_valid_pdf_bytes():
    docx_bytes = _make_docx("Hello world", "Second paragraph")
    pdf_bytes = format_converter.convert_to_pdf(docx_bytes)

    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 0
    assert pdf_bytes[:5] == b"%PDF-"


def test_convert_to_pdf_produces_nonempty_valid_pdf():
    docx_bytes = _make_docx("[Email 1] contacted Person 1")
    pdf_bytes = format_converter.convert_to_pdf(docx_bytes)

    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes[:5] == b"%PDF-"
    assert len(pdf_bytes) > 100  # Non-trivial PDF content


def test_convert_to_pdf_handles_headings():
    docx_bytes = _make_docx(headings=[(1, "Title"), (2, "Subtitle")])
    pdf_bytes = format_converter.convert_to_pdf(docx_bytes)

    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes[:5] == b"%PDF-"


def test_convert_to_markdown_returns_valid_markdown():
    docx_bytes = _make_docx("Hello world", "Second paragraph")
    md = format_converter.convert_to_markdown(docx_bytes)

    assert isinstance(md, str)
    assert "Hello world" in md
    assert "Second paragraph" in md


def test_convert_to_markdown_preserves_headings():
    docx_bytes = _make_docx(
        "Body text",
        headings=[(1, "Main Title"), (2, "Sub Title"), (3, "Section")],
    )
    md = format_converter.convert_to_markdown(docx_bytes)

    assert "# Main Title" in md
    assert "## Sub Title" in md
    assert "### Section" in md


def test_convert_to_markdown_preserves_redacted_text():
    docx_bytes = _make_docx(
        "[Email 1] contacted Person 1 at X Private Limited."
    )
    md = format_converter.convert_to_markdown(docx_bytes)

    assert "[Email 1]" in md
    assert "Person 1" in md
    assert "X Private Limited" in md


# ── API endpoint tests ───────────────────────────────────────

def _fake_process_document(_content: bytes):
    doc = DocxDocument()
    doc.add_paragraph("Person 1 works at X Private Limited.")
    out = BytesIO()
    doc.save(out)
    return (
        {"INDIVIDUAL": 1, "PRIVATE_LIMITED": 1},
        2,
        out.getvalue(),
        [
            {"category": "INDIVIDUAL", "placeholder": "Person 1", "location": "Paragraph 1", "original_text": "Alice"},
            {"category": "PRIVATE_LIMITED", "placeholder": "X Private Limited", "location": "Paragraph 1", "original_text": "Acme Private Limited"},
        ],
        [],
    )


def test_download_default_format_is_docx(monkeypatch):
    monkeypatch.setattr(server.nlp_engine, "initialize_models", lambda: None)
    monkeypatch.setattr(server, "process_document", _fake_process_document)

    client = TestClient(server.app)
    resp = client.post(
        "/api/redact",
        files={"file": ("test.docx", b"fake", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 200
    file_id = resp.json()["file_id"]

    dl = client.get(f"/api/download/{file_id}")
    assert dl.status_code == 200
    assert dl.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert dl.headers["content-disposition"].endswith('.docx"')


def test_download_pdf_format(monkeypatch):
    monkeypatch.setattr(server.nlp_engine, "initialize_models", lambda: None)
    monkeypatch.setattr(server, "process_document", _fake_process_document)

    client = TestClient(server.app)
    resp = client.post(
        "/api/redact",
        files={"file": ("test.docx", b"fake", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    file_id = resp.json()["file_id"]

    dl = client.get(f"/api/download/{file_id}?output_format=pdf")
    assert dl.status_code == 200
    assert dl.headers["content-type"] == "application/pdf"
    assert dl.content[:5] == b"%PDF-"
    assert dl.headers["content-disposition"].endswith('.pdf"')


def test_download_markdown_format(monkeypatch):
    monkeypatch.setattr(server.nlp_engine, "initialize_models", lambda: None)
    monkeypatch.setattr(server, "process_document", _fake_process_document)

    client = TestClient(server.app)
    resp = client.post(
        "/api/redact",
        files={"file": ("test.docx", b"fake", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    file_id = resp.json()["file_id"]

    dl = client.get(f"/api/download/{file_id}?output_format=md")
    assert dl.status_code == 200
    assert "text/markdown" in dl.headers["content-type"]
    assert "Person 1" in dl.text
    assert "X Private Limited" in dl.text
    assert dl.headers["content-disposition"].endswith('.md"')


def test_download_docx_explicit_format(monkeypatch):
    monkeypatch.setattr(server.nlp_engine, "initialize_models", lambda: None)
    monkeypatch.setattr(server, "process_document", _fake_process_document)

    client = TestClient(server.app)
    resp = client.post(
        "/api/redact",
        files={"file": ("test.docx", b"fake", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    file_id = resp.json()["file_id"]

    dl = client.get(f"/api/download/{file_id}?output_format=docx")
    assert dl.status_code == 200
    assert dl.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def test_all_formats_contain_same_redacted_text(monkeypatch):
    """All output formats must contain the same redacted text."""
    monkeypatch.setattr(server.nlp_engine, "initialize_models", lambda: None)
    monkeypatch.setattr(server, "process_document", _fake_process_document)

    client = TestClient(server.app)
    resp = client.post(
        "/api/redact",
        files={"file": ("test.docx", b"fake", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    file_id = resp.json()["file_id"]

    # DOCX
    dl_docx = client.get(f"/api/download/{file_id}?output_format=docx")
    doc = DocxDocument(BytesIO(dl_docx.content))
    docx_text = "\n".join(p.text for p in doc.paragraphs)

    # PDF — verify it's a valid PDF (text is compressed, can't check raw bytes)
    dl_pdf = client.get(f"/api/download/{file_id}?output_format=pdf")
    assert dl_pdf.content[:5] == b"%PDF-"

    # Markdown
    dl_md = client.get(f"/api/download/{file_id}?output_format=md")
    md_text = dl_md.text

    # DOCX and Markdown should contain the same redacted text
    assert "Person 1" in docx_text
    assert "Person 1" in md_text
    assert "X Private Limited" in docx_text
    assert "X Private Limited" in md_text
