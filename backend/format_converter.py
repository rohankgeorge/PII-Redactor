"""Convert redacted DOCX documents to PDF or Markdown format."""

from __future__ import annotations

import io
from docx import Document as DocxDocument
from fpdf import FPDF


def convert_to_pdf(docx_bytes: bytes) -> bytes:
    """Extract text from a DOCX and generate a PDF using fpdf2.

    Preserves paragraph structure and heading levels.
    Returns raw PDF bytes.
    """
    doc = DocxDocument(io.BytesIO(docx_bytes))
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # Use built-in Helvetica (no external font files needed).
    for para in doc.paragraphs:
        text = para.text
        style_name = (para.style.name or "").lower()

        if style_name.startswith("heading 1"):
            pdf.set_font("Helvetica", "B", 16)
            pdf.multi_cell(0, 8, text)
            pdf.ln(4)
        elif style_name.startswith("heading 2"):
            pdf.set_font("Helvetica", "B", 14)
            pdf.multi_cell(0, 7, text)
            pdf.ln(3)
        elif style_name.startswith("heading"):
            pdf.set_font("Helvetica", "B", 12)
            pdf.multi_cell(0, 6, text)
            pdf.ln(2)
        else:
            pdf.set_font("Helvetica", "", 11)
            if text.strip():
                pdf.multi_cell(0, 6, text)
                pdf.ln(2)
            else:
                pdf.ln(4)

    # Also include table content.
    for table in doc.tables:
        for row in table.rows:
            cells_text = [cell.text.strip() for cell in row.cells]
            line = " | ".join(cells_text)
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(0, 5, line)
        pdf.ln(3)

    return bytes(pdf.output())


def convert_to_markdown(docx_bytes: bytes) -> str:
    """Extract paragraphs from a DOCX and convert to Markdown syntax.

    Preserves headings and paragraph structure.
    Returns a Markdown string.
    """
    doc = DocxDocument(io.BytesIO(docx_bytes))
    lines: list[str] = []

    for para in doc.paragraphs:
        text = para.text
        style_name = (para.style.name or "").lower()

        if style_name.startswith("heading 1"):
            lines.append(f"# {text}")
            lines.append("")
        elif style_name.startswith("heading 2"):
            lines.append(f"## {text}")
            lines.append("")
        elif style_name.startswith("heading 3"):
            lines.append(f"### {text}")
            lines.append("")
        elif style_name.startswith("heading"):
            # Heading 4+ → ####
            level = 4
            for ch in style_name:
                if ch.isdigit():
                    level = min(int(ch), 6)
                    break
            lines.append(f"{'#' * level} {text}")
            lines.append("")
        elif text.strip():
            lines.append(text)
            lines.append("")
        else:
            lines.append("")

    # Tables → markdown table syntax.
    for table in doc.tables:
        if not table.rows:
            continue
        header_cells = [cell.text.strip() for cell in table.rows[0].cells]
        lines.append("| " + " | ".join(header_cells) + " |")
        lines.append("| " + " | ".join("---" for _ in header_cells) + " |")
        for row in table.rows[1:]:
            cells = [cell.text.strip() for cell in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
        lines.append("")

    return "\n".join(lines)
