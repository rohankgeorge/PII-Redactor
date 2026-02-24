"""Tests for Milestone 3: force_redact_terms and placeholder_overrides in apply-redaction."""

from io import BytesIO

from docx import Document as DocxDocument
from fastapi.testclient import TestClient

import server


def _fake_process_document_m3(_content: bytes):
    """Fake processor that leaves 'SecretProject' un-redacted in the output."""
    doc = DocxDocument()
    doc.add_paragraph("Person 1 works at X Private Limited on SecretProject.")
    out = BytesIO()
    doc.save(out)
    return (
        {"INDIVIDUAL": 1, "PRIVATE_LIMITED": 1},
        2,
        out.getvalue(),
        [
            {"category": "INDIVIDUAL", "placeholder": "Person 1", "location": "Paragraph 1", "original_text": "Alice"},
            {"category": "PRIVATE_LIMITED", "placeholder": "X Private Limited", "location": "Paragraph 1", "original_text": "Acme Private Limited"},
        ],
        [],
    )


def _analyze(monkeypatch, fake_fn=None):
    """Helper: run analyze and return (client, analysis_json)."""
    monkeypatch.setattr(server.nlp_engine, "initialize_models", lambda: None)
    monkeypatch.setattr(server, "process_document", fake_fn or _fake_process_document_m3)
    client = TestClient(server.app)
    resp = client.post(
        "/api/analyze",
        files={"file": ("test.docx", b"fake", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 200
    return client, resp.json()


def _download_docx_text(client, file_id):
    """Helper: download a file and return the concatenated paragraph text."""
    dl = client.get(f"/api/download/{file_id}")
    assert dl.status_code == 200
    doc = DocxDocument(BytesIO(dl.content))
    return "\n".join(p.text for p in doc.paragraphs)


# ── force_redact_terms ──────────────────────────────────────

def test_force_redact_terms_redacts_additional_term(monkeypatch):
    """A term not detected by the engine gets redacted when specified as a force-redact term."""
    client, analysis = _analyze(monkeypatch)
    applied = client.post("/api/apply-redaction", json={
        "analysis_id": analysis["analysis_id"],
        "force_redact_terms": ["SecretProject"],
    })
    assert applied.status_code == 200
    payload = applied.json()

    # Output DOCX must not contain the force-redacted term
    text = _download_docx_text(client, payload["file_id"])
    assert "SecretProject" not in text

    # Audit log should have an entry for the force-redacted term
    force_entries = [r for r in payload["audit_log"] if r.get("original_text") == "SecretProject"]
    assert len(force_entries) >= 1

    # Total and stats should reflect the additional redaction
    assert payload["total"] >= 3  # 2 original + 1 force


def test_force_redact_terms_overlap_existing_no_duplicate(monkeypatch):
    """Force-redacting a term whose original text is already replaced does not cause errors."""
    client, analysis = _analyze(monkeypatch)
    # "Alice" was already replaced by "Person 1" in the DOCX — force-redacting "Alice" is a no-op.
    applied = client.post("/api/apply-redaction", json={
        "analysis_id": analysis["analysis_id"],
        "force_redact_terms": ["Alice"],
    })
    assert applied.status_code == 200
    payload = applied.json()
    # Original 2 redactions should remain
    assert payload["total"] >= 2


def test_force_redact_empty_list_is_noop(monkeypatch):
    """An empty force_redact_terms list changes nothing."""
    client, analysis = _analyze(monkeypatch)
    applied = client.post("/api/apply-redaction", json={
        "analysis_id": analysis["analysis_id"],
        "force_redact_terms": [],
    })
    assert applied.status_code == 200
    assert applied.json()["total"] == 2


# ── placeholder_overrides ───────────────────────────────────

def test_placeholder_overrides_replaces_auto_placeholder(monkeypatch):
    """A custom override replaces the auto-generated placeholder in the DOCX."""
    client, analysis = _analyze(monkeypatch)
    # Find the INDIVIDUAL candidate
    ind_cand = next(c for c in analysis["candidates"] if c["category"] == "INDIVIDUAL")

    applied = client.post("/api/apply-redaction", json={
        "analysis_id": analysis["analysis_id"],
        "placeholder_overrides": {ind_cand["candidate_id"]: "THE_CEO"},
    })
    assert applied.status_code == 200
    payload = applied.json()

    text = _download_docx_text(client, payload["file_id"])
    assert "THE_CEO" in text
    assert "Person 1" not in text  # auto-placeholder should be replaced


def test_placeholder_overrides_unknown_candidate_ignored(monkeypatch):
    """Unknown candidate IDs in overrides are silently ignored."""
    client, analysis = _analyze(monkeypatch)
    applied = client.post("/api/apply-redaction", json={
        "analysis_id": analysis["analysis_id"],
        "placeholder_overrides": {"cand_99999": "CUSTOM"},
    })
    assert applied.status_code == 200
    assert applied.json()["total"] == 2  # unchanged


def test_placeholder_overrides_updates_audit_log(monkeypatch):
    """When a placeholder is overridden, the audit log reflects the custom text."""
    client, analysis = _analyze(monkeypatch)
    ind_cand = next(c for c in analysis["candidates"] if c["category"] == "INDIVIDUAL")

    applied = client.post("/api/apply-redaction", json={
        "analysis_id": analysis["analysis_id"],
        "placeholder_overrides": {ind_cand["candidate_id"]: "THE_CEO"},
    })
    assert applied.status_code == 200
    payload = applied.json()

    ind_rows = [r for r in payload["audit_log"] if r["category"] == "INDIVIDUAL"]
    assert any(r["placeholder"] == "THE_CEO" for r in ind_rows)


# ── Combined ────────────────────────────────────────────────

def test_force_redact_and_overrides_combined(monkeypatch):
    """Both features work together in a single request."""
    client, analysis = _analyze(monkeypatch)
    ind_cand = next(c for c in analysis["candidates"] if c["category"] == "INDIVIDUAL")

    applied = client.post("/api/apply-redaction", json={
        "analysis_id": analysis["analysis_id"],
        "force_redact_terms": ["SecretProject"],
        "placeholder_overrides": {ind_cand["candidate_id"]: "THE_CEO"},
    })
    assert applied.status_code == 200
    payload = applied.json()

    text = _download_docx_text(client, payload["file_id"])
    assert "THE_CEO" in text
    assert "Person 1" not in text
    assert "SecretProject" not in text
    assert payload["total"] >= 3
